"""
Templates router - Provides sample Excel templates and format guides.
"""

from fastapi import APIRouter, Response
from fastapi.responses import StreamingResponse
from typing import Dict, Any
import io
import logging

from app.services.ai_analyzer import AIAnalyzer
from app.models.upload import DataType

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/templates", tags=["Templates"])


# Sample data for each data type
SAMPLE_DATA = {
    "sales": [
        {"date": "2024-12-01", "amount": 1500, "sku": "COFFEE-LATTE", "quantity": 10, "outlet": "Main Store"},
        {"date": "2024-12-01", "amount": 2200, "sku": "TEA-MASALA", "quantity": 15, "outlet": "Main Store"},
        {"date": "2024-12-01", "amount": 850, "sku": "SANDWICH-VEG", "quantity": 5, "outlet": "Main Store"},
        {"date": "2024-12-02", "amount": 1800, "sku": "COFFEE-ESPRESSO", "quantity": 12, "outlet": "Main Store"},
        {"date": "2024-12-02", "amount": 1200, "sku": "PASTRY-CHOCO", "quantity": 8, "outlet": "Branch 1"},
        {"date": "2024-12-02", "amount": 3500, "sku": "MEAL-COMBO", "quantity": 7, "outlet": "Main Store"},
        {"date": "2024-12-03", "amount": 950, "sku": "JUICE-ORANGE", "quantity": 6, "outlet": "Branch 1"},
        {"date": "2024-12-03", "amount": 2800, "sku": "PIZZA-MARGHERITA", "quantity": 4, "outlet": "Main Store"},
        {"date": "2024-12-03", "amount": 1650, "sku": "COFFEE-LATTE", "quantity": 11, "outlet": "Branch 1"},
        {"date": "2024-12-04", "amount": 4200, "sku": "MEAL-PREMIUM", "quantity": 6, "outlet": "Main Store"},
    ],
    "wastage": [
        {"date": "2024-12-01", "sku": "MILK-FULL", "quantity": 5, "value": 250, "reason": "Expired"},
        {"date": "2024-12-01", "sku": "BREAD-WHITE", "quantity": 10, "value": 150, "reason": "Stale"},
        {"date": "2024-12-02", "sku": "FRUIT-BANANA", "quantity": 8, "value": 120, "reason": "Overripe"},
        {"date": "2024-12-02", "sku": "SALAD-GREEN", "quantity": 3, "value": 180, "reason": "Wilted"},
        {"date": "2024-12-03", "sku": "YOGURT-PLAIN", "quantity": 4, "value": 200, "reason": "Expired"},
        {"date": "2024-12-03", "sku": "CREAM-WHIP", "quantity": 2, "value": 160, "reason": "Spoiled"},
        {"date": "2024-12-04", "sku": "SANDWICH-VEG", "quantity": 3, "value": 225, "reason": "Not sold"},
        {"date": "2024-12-04", "sku": "PASTRY-FRUIT", "quantity": 5, "value": 350, "reason": "Stale"},
    ],
    "purchase": [
        {"date": "2024-12-01", "supplier": "Fresh Dairy Ltd", "sku": "MILK-BULK", "quantity": 100, "value": 4500},
        {"date": "2024-12-01", "supplier": "Green Farms", "sku": "VEG-ASSORTED", "quantity": 50, "value": 3000},
        {"date": "2024-12-02", "supplier": "Coffee Beans Co", "sku": "COFFEE-ARABICA", "quantity": 25, "value": 12500},
        {"date": "2024-12-02", "supplier": "Bakery Supplies", "sku": "FLOUR-MAIDA", "quantity": 40, "value": 2000},
        {"date": "2024-12-03", "supplier": "Fresh Dairy Ltd", "sku": "CREAM-FRESH", "quantity": 20, "value": 1800},
        {"date": "2024-12-03", "supplier": "Fruit Mart", "sku": "FRUIT-MIX", "quantity": 30, "value": 2400},
        {"date": "2024-12-04", "supplier": "Green Farms", "sku": "HERBS-FRESH", "quantity": 15, "value": 750},
    ],
    "inventory": [
        {"date": "2024-12-01", "sku": "COFFEE-ARABICA", "opening_stock": 100, "closing_stock": 85, "movement": -15},
        {"date": "2024-12-01", "sku": "MILK-FULL", "opening_stock": 50, "closing_stock": 35, "movement": -15},
        {"date": "2024-12-01", "sku": "SUGAR-WHITE", "opening_stock": 40, "closing_stock": 38, "movement": -2},
        {"date": "2024-12-02", "sku": "COFFEE-ARABICA", "opening_stock": 85, "closing_stock": 72, "movement": -13},
        {"date": "2024-12-02", "sku": "MILK-FULL", "opening_stock": 35, "closing_stock": 80, "movement": 45},
        {"date": "2024-12-02", "sku": "BREAD-WHITE", "opening_stock": 25, "closing_stock": 50, "movement": 25},
        {"date": "2024-12-03", "sku": "COFFEE-ARABICA", "opening_stock": 72, "closing_stock": 60, "movement": -12},
        {"date": "2024-12-03", "sku": "CREAM-FRESH", "opening_stock": 10, "closing_stock": 25, "movement": 15},
    ],
    "staff": [
        {"date": "2024-12-01", "staff_name": "Rahul Kumar", "type": "attendance", "description": "On time", "outlet": "Main Store"},
        {"date": "2024-12-01", "staff_name": "Priya Sharma", "type": "attendance", "description": "On time", "outlet": "Main Store"},
        {"date": "2024-12-01", "staff_name": "Amit Singh", "type": "late", "description": "30 min late", "outlet": "Branch 1"},
        {"date": "2024-12-02", "staff_name": "Rahul Kumar", "type": "issue", "description": "Customer complaint handled well", "outlet": "Main Store"},
        {"date": "2024-12-02", "staff_name": "Neha Gupta", "type": "attendance", "description": "On time", "outlet": "Main Store"},
        {"date": "2024-12-02", "staff_name": "Priya Sharma", "type": "performance", "description": "Highest sales of the day", "outlet": "Main Store"},
        {"date": "2024-12-03", "staff_name": "Amit Singh", "type": "attendance", "description": "On time", "outlet": "Branch 1"},
        {"date": "2024-12-03", "staff_name": "Vikash Patel", "type": "absent", "description": "Sick leave", "outlet": "Branch 1"},
    ]
}


@router.get("/download/{data_type}")
async def download_sample_excel(data_type: str):
    """
    Download a sample Excel file for the specified data type.
    This serves as a reference template for users.
    """
    try:
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils.dataframe import dataframe_to_rows
        
        if data_type not in SAMPLE_DATA:
            return {"error": f"Unknown data type: {data_type}"}
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"{data_type.title()} Data"
        
        # Get sample data
        data = SAMPLE_DATA[data_type]
        df = pd.DataFrame(data)
        
        # Style definitions
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Add instructions row
        ws.merge_cells('A1:F1')
        ws['A1'] = f"📋 Sample {data_type.title()} Data Template - Use this format for your data"
        ws['A1'].font = Font(bold=True, size=12, color="1E3A8A")
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Add empty row
        ws.append([])
        
        # Add headers
        headers = list(df.columns)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # Add data rows
        for row_idx, row in enumerate(df.values, start=4):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='left')
        
        # Adjust column widths
        for col_idx, header in enumerate(headers, start=1):
            max_length = max(len(str(header)), max(len(str(val)) for val in df[header]))
            ws.column_dimensions[chr(64 + col_idx)].width = min(max_length + 2, 30)
        
        # Add format notes at the bottom
        notes_row = len(data) + 6
        ws.merge_cells(f'A{notes_row}:F{notes_row}')
        ws[f'A{notes_row}'] = "📝 Format Notes:"
        ws[f'A{notes_row}'].font = Font(bold=True, size=11)
        
        format_notes = get_format_notes(data_type)
        for i, note in enumerate(format_notes, start=1):
            ws[f'A{notes_row + i}'] = f"  • {note}"
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"sample_{data_type}_template.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )
        
    except ImportError as e:
        logger.error(f"Missing dependency: {e}")
        return {"error": "Excel generation requires pandas and openpyxl. Please install them."}
    except Exception as e:
        logger.error(f"Error generating Excel: {e}")
        return {"error": str(e)}


@router.get("/format-guide/{data_type}")
async def get_format_guide(data_type: str) -> Dict[str, Any]:
    """
    Get AI-powered format guide for a specific data type.
    Includes required columns, examples, and tips.
    """
    analyzer = AIAnalyzer()
    guide = await analyzer.generate_excel_format_guide(data_type)
    return guide


@router.get("/sample-data/{data_type}")
async def get_sample_data(data_type: str) -> Dict[str, Any]:
    """
    Get sample data as JSON for preview or reference.
    """
    if data_type not in SAMPLE_DATA:
        return {"error": f"Unknown data type: {data_type}"}
    
    return {
        "data_type": data_type,
        "sample_count": len(SAMPLE_DATA[data_type]),
        "columns": list(SAMPLE_DATA[data_type][0].keys()),
        "data": SAMPLE_DATA[data_type],
        "format_notes": get_format_notes(data_type)
    }


@router.get("/all-types")
async def get_all_data_types() -> Dict[str, Any]:
    """
    Get information about all supported data types.
    """
    return {
        "data_types": [
            {
                "type": "sales",
                "display_name": "Sales Transactions",
                "description": "Daily sales data with amounts, SKUs, and quantities",
                "required_columns": ["date", "amount"],
                "optional_columns": ["sku", "quantity", "outlet"],
                "icon": "💰"
            },
            {
                "type": "wastage",
                "display_name": "Wastage Records",
                "description": "Track wasted items, quantities, and reasons",
                "required_columns": ["date", "sku", "quantity"],
                "optional_columns": ["value", "reason", "outlet"],
                "icon": "🗑️"
            },
            {
                "type": "purchase",
                "display_name": "Purchase/Procurement",
                "description": "Supplier purchases and procurement data",
                "required_columns": ["date", "value"],
                "optional_columns": ["supplier", "sku", "quantity", "outlet"],
                "icon": "📦"
            },
            {
                "type": "inventory",
                "display_name": "Inventory Tracking",
                "description": "Daily stock levels and movements",
                "required_columns": ["date", "sku"],
                "optional_columns": ["opening_stock", "closing_stock", "movement", "outlet"],
                "icon": "📊"
            },
            {
                "type": "staff",
                "display_name": "Staff Logs",
                "description": "Attendance, issues, and performance logs",
                "required_columns": ["date", "staff_name", "type"],
                "optional_columns": ["description", "outlet"],
                "icon": "👥"
            }
        ]
    }


def get_format_notes(data_type: str) -> list:
    """Get format notes for a data type."""
    notes = {
        "sales": [
            "Date format: YYYY-MM-DD (e.g., 2024-12-15)",
            "Amount: Plain number without currency symbols (e.g., 1500, not ₹1,500)",
            "SKU: Consistent product codes (e.g., COFFEE-LATTE)",
            "Quantity: Whole numbers only"
        ],
        "wastage": [
            "Date format: YYYY-MM-DD",
            "SKU: Same codes used in sales data",
            "Quantity: Amount wasted (can be decimal for weight-based items)",
            "Value: Cost of wasted items (optional but recommended)",
            "Reason: Brief description (Expired, Stale, Damaged, etc.)"
        ],
        "purchase": [
            "Date format: YYYY-MM-DD",
            "Value: Total purchase value for the line item",
            "Supplier: Consistent supplier names",
            "SKU: Product codes matching your inventory"
        ],
        "inventory": [
            "Date format: YYYY-MM-DD",
            "SKU: Consistent product codes",
            "Opening/Closing Stock: Stock counts at start/end of day",
            "Movement: Automatically calculated as closing - opening"
        ],
        "staff": [
            "Date format: YYYY-MM-DD",
            "Staff Name: Full name, consistent spelling",
            "Type: attendance, late, absent, issue, performance",
            "Description: Brief notes about the log entry"
        ]
    }
    return notes.get(data_type, ["Use YYYY-MM-DD date format", "Use plain numbers without symbols"])
